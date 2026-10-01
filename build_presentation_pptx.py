import math
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from scipy.integrate import odeint


BASE_DIR = Path("/Users/bandiyaa/Desktop/4090")
ASSETS_DIR = BASE_DIR / "presentation_assets"
PPTX_PATH = BASE_DIR / "project_presentation_ready.pptx"
NOTES_PATH = BASE_DIR / "project_presentation_speaker_notes.md"

M = 2


def model(y, t, N, beta, gamma, rho, sigma):
    I = y[: M + 1]
    S = y[M + 1]
    lam = beta * S / N

    dI = np.zeros(M + 1)
    dI[0] = -lam * I[0]
    for i in range(1, M):
        dI[i] = (1 - rho[i - 1] - sigma[i - 1]) * lam * I[i - 1] - lam * I[i]
    dI[M] = (1 - rho[M - 1] - sigma[M - 1]) * lam * I[M - 1] - (rho[M] + sigma[M]) * lam * I[M]

    dS = sum(rho[i] * lam * I[i] for i in range(M + 1)) - gamma * S
    return np.concatenate([dI, [dS]])


def run_simulation(N, beta, gamma, rho, sigma, I0_init, T=120, npts=600):
    y0 = np.zeros(M + 2)
    y0[0] = I0_init
    y0[M + 1] = N - I0_init
    t = np.linspace(0, T, npts)
    sol = odeint(model, y0, t, args=(N, beta, gamma, rho, sigma))
    I_all = sol[:, : M + 1]
    S = sol[:, M + 1]
    R = N - I_all.sum(axis=1) - S
    return t, I_all, S, R


def ensure_assets():
    ASSETS_DIR.mkdir(exist_ok=True)
    plt.style.use("seaborn-v0_8-whitegrid")

    N = 1000.0
    T = 120
    npts = 600
    rho_viral = np.array([0.05, 0.15, 0.25])
    sigma_viral = np.array([0.01, 0.02, 0.04])
    beta_viral = 1.5
    gamma_viral = 0.08

    rho_mod = np.array([0.05, 0.15, 0.25])
    sigma_mod = np.array([0.01, 0.02, 0.04])
    beta_mod = 1.13
    gamma_mod = 0.089

    rho_fizzle = np.array([0.01, 0.02, 0.03])
    sigma_fizzle = np.array([0.10, 0.15, 0.20])
    beta_fizzle = 0.15
    gamma_fizzle = 0.6

    t1, I1, S1, R1 = run_simulation(N, beta_viral, gamma_viral, rho_viral, sigma_viral, N - 15, T, npts)
    t2, I2, S2, R2 = run_simulation(N, beta_fizzle, gamma_fizzle, rho_fizzle, sigma_fizzle, N - 15, T, npts)
    t3, I3, S3, R3 = run_simulation(N, beta_mod, gamma_mod, rho_mod, sigma_mod, N - 15, T, npts)

    comparison = ASSETS_DIR / "comparison.png"
    fig, ax = plt.subplots(figsize=(10, 5.4))
    ax.plot(t1, S1, color="#1f77b4", linewidth=3, label="Viral")
    ax.plot(t3, S3, color="#7f7f7f", linewidth=3, label="Moderate")
    ax.plot(t2, S2, color="#d62728", linewidth=3, label="Fizzle")
    ax.set_xlabel("Time")
    ax.set_ylabel("Active sharers S(t)")
    ax.set_title("Sharer Dynamics Across Scenarios", fontsize=18, pad=14)
    ax.legend(frameon=True)
    fig.tight_layout()
    fig.savefig(comparison, dpi=220, bbox_inches="tight")
    plt.close(fig)

    exposure = ASSETS_DIR / "exposure_classes.png"
    fig, ax = plt.subplots(figsize=(10, 5.4))
    ax.plot(t1, I1[:, 0], linewidth=2.5, label="I0")
    ax.plot(t1, I1[:, 1], linewidth=2.5, label="I1")
    ax.plot(t1, I1[:, 2], linewidth=2.5, label="I2")
    ax.plot(t1, S1, linewidth=2.8, label="S")
    ax.plot(t1, R1, linewidth=2.8, label="R")
    ax.set_xlabel("Time")
    ax.set_ylabel("Users")
    ax.set_title("Viral Scenario: Exposure Classes and Outcomes", fontsize=18, pad=14)
    ax.legend(ncol=5, frameon=True, loc="upper center")
    fig.tight_layout()
    fig.savefig(exposure, dpi=220, bbox_inches="tight")
    plt.close(fig)

    sensitivity = ASSETS_DIR / "sensitivity_gamma.png"
    gammas = np.linspace(0.06, 0.35, 25)
    beta_sens = 1.2
    peaks_g = []
    for g in gammas:
        t, _, S, _ = run_simulation(N, beta_sens, g, rho_viral, sigma_viral, N - 15, T, npts)
        peaks_g.append(float(S.max()))
    fig, ax = plt.subplots(figsize=(10, 5.4))
    ax.plot(gammas, peaks_g, color="#ff7f0e", linewidth=3)
    ax.set_xlabel("Gamma")
    ax.set_ylabel("Peak active sharers")
    ax.set_title("Sensitivity of Peak Sharing to Gamma", fontsize=18, pad=14)
    fig.tight_layout()
    fig.savefig(sensitivity, dpi=220, bbox_inches="tight")
    plt.close(fig)

    return {
        "comparison": comparison,
        "exposure": exposure,
        "sensitivity": sensitivity,
        "stats": {
            "viral_peak": float(S1.max()),
            "viral_peak_t": float(t1[S1.argmax()]),
            "viral_R": float(R1[-1]),
            "moderate_peak": float(S3.max()),
            "moderate_peak_t": float(t3[S3.argmax()]),
            "moderate_R": float(R3[-1]),
            "fizzle_peak": float(S2.max()),
            "fizzle_peak_t": float(t2[S2.argmax()]),
            "fizzle_R": float(R2[-1]),
        },
    }


def add_bg(slide, color=(246, 248, 252)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*color)


def add_header_bar(slide, title, subtitle=None):
    left, top, width, height = Inches(0), Inches(0), Inches(13.333), Inches(0.75)
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(34, 64, 121)
    shape.line.fill.background()

    text = slide.shapes.add_textbox(Inches(0.5), Inches(0.14), Inches(10.8), Inches(0.35))
    p = text.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.5), Inches(0.43), Inches(9.5), Inches(0.22))
        p2 = sub.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(222, 230, 245)


def add_footer(slide, text="MATH 4090 Project Presentation"):
    box = slide.shapes.add_textbox(Inches(0.45), Inches(7.05), Inches(6.5), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(90, 98, 110)


def add_bullets(slide, items, left=0.7, top=1.2, width=5.8, height=4.8, font_size=22, color=(32, 39, 49)):
    tx = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tx.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(12)
    return tx


def add_note_box(slide, title, body, left=7.0, top=1.4, width=5.6, height=2.6, fill=(232, 240, 254)):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*fill)
    shape.line.color.rgb = RGBColor(163, 188, 237)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(34, 64, 121)
    p.space_after = Pt(8)
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(32, 39, 49)
    return shape


def add_flow_diagram(slide):
    labels = ["I0\nNot yet exposed", "I1\nOne exposure", "I2\nTwo exposures", "S\nSharers", "R\nSaturated"]
    colors = [(220, 234, 255), (220, 234, 255), (220, 234, 255), (222, 245, 229), (252, 232, 230)]
    x_positions = [0.7, 3.0, 5.3, 7.8, 10.2]
    for i, (label, color, x) in enumerate(zip(labels, colors, x_positions)):
        shp = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.2), Inches(1.7), Inches(1.2))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(*color)
        shp.line.color.rgb = RGBColor(110, 125, 150)
        tf = shp.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.color.rgb = RGBColor(32, 39, 49)
        if i < len(labels) - 1:
            arr = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.CHEVRON, Inches(x + 1.8), Inches(2.47), Inches(0.55), Inches(0.5))
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(160, 177, 201)
            arr.line.fill.background()


def add_table_slide(prs, title, subtitle, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, title, subtitle)
    add_footer(slide)

    table = slide.shapes.add_table(len(rows), 4, Inches(0.7), Inches(1.45), Inches(11.9), Inches(3.0)).table
    widths = [2.2, 2.2, 2.6, 2.8]
    for i, w in enumerate(widths):
        table.columns[i].width = Inches(w)

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(17 if r_idx == 0 else 16)
                p.font.bold = r_idx == 0
                p.font.color.rgb = RGBColor(32, 39, 49)
            if r_idx == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(221, 232, 250)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(250, 251, 253)

    add_note_box(
        slide,
        "How to say this clearly",
        "Viral has the highest peak, moderate is smaller and slower, and fizzle stays near the initial seed. This is the main numerical message.",
        left=0.9,
        top=5.0,
        width=11.5,
        height=1.35,
    )
    return slide


def build_ppt():
    assets = ensure_assets()
    stats = assets["stats"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_side_panel(slide, title, body, top=1.45, height=4.9, fill=(248, 250, 253)):
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(8.9),
            Inches(top),
            Inches(3.8),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*fill)
        shape.line.color.rgb = RGBColor(210, 218, 230)
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(34, 64, 121)
        p.space_after = Pt(10)
        for line in body:
            para = tf.add_paragraph()
            para.text = line
            para.font.size = Pt(15)
            para.font.color.rgb = RGBColor(32, 39, 49)
            para.space_after = Pt(8)

    def add_caption(slide, text, left=0.85, top=6.45, width=7.8):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.4))
        p = box.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(11)
        p.font.italic = True
        p.font.color.rgb = RGBColor(90, 98, 110)

    def add_small_table(slide, rows, left=0.8, top=1.65, width=7.55, height=3.2):
        table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)).table
        for c in range(len(rows[0])):
            table.columns[c].width = Inches(width / len(rows[0]))
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.text = val
                for p in cell.text_frame.paragraphs:
                    p.font.size = Pt(16 if r_idx == 0 else 15)
                    p.font.bold = r_idx == 0
                    p.font.color.rgb = RGBColor(32, 39, 49)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(*(225, 235, 250) if r_idx == 0 else (250, 251, 253))
        return table

    # Slide 1: title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, (244, 247, 252))
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(27, 54, 104)
    band.line.fill.background()
    add_footer(slide, "MATH 4090 | Project presentation")
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(8.4), Inches(1.9))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Multi-Exposure Modelling of Information Spread Through Social Media Reposting"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(25, 31, 40)
    p.space_after = Pt(14)
    p2 = tf.add_paragraph()
    p2.text = "Lavanya Kumar"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(70, 78, 90)
    p3 = tf.add_paragraph()
    p3.text = "Designed for a 12-minute talk plus Q&A"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(90, 98, 110)
    add_side_panel(
        slide,
        "Presentation focus",
        [
            "Why repeated exposure matters",
            "How the ODE model is built",
            "What the simulations show",
            "What remains for the final report",
        ],
        top=1.55,
        height=2.85,
        fill=(232, 240, 254),
    )
    add_flow_diagram(slide)

    # Slide 2: background
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "1. Motivation and Background", "Why this problem is interesting and where it fits in the literature")
    add_footer(slide)
    add_bullets(
        slide,
        [
            "On platforms such as Instagram, users often encounter the same content more than once before deciding whether to repost it.",
            "That means information spread can depend on accumulated exposure, not only on a one-time contact.",
            "Classical rumor and epidemic-style models describe spreaders and non-spreaders, but they usually do not track repeated exposure explicitly.",
            "This project adds a multi-exposure structure while keeping the model mathematically tractable.",
        ],
        width=7.8,
        font_size=20,
    )
    add_side_panel(
        slide,
        "Related modelling",
        [
            "Rumor models: spreaders and stiflers",
            "Threshold models: adoption after sufficient influence",
            "Complex contagion: reinforcement through repeated exposure",
            "My model: ODE framework with explicit exposure classes",
        ],
        fill=(240, 244, 248),
    )

    # Slide 3: compartments and assumptions
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "2. Compartments, Parameters, and Assumptions", "Model ingredients stated before the equations")
    add_footer(slide)
    rows = [
        ["Symbol", "Meaning", "Where it appears"],
        ["I0, I1, I2", "Exposure classes before sharing", "State variables"],
        ["S", "Active sharers", "Drives lambda"],
        ["R", "Saturated users", "Collects users leaving spread"],
        ["beta", "Repost intensity / visibility", "lambda = beta S / N"],
        ["gamma", "Rate sharers stop sharing", "dS/dt and dR/dt"],
        ["rho_i", "Probability of sharing at exposure i", "Inflow into S"],
        ["sigma_i", "Probability of saturation at exposure i", "Inflow into R"],
    ]
    add_small_table(slide, rows, width=7.9, height=4.25)
    add_side_panel(
        slide,
        "Assumptions",
        [
            "Population is fixed: I0 + I1 + I2 + S + R = N",
            "Exposure rate is mean-field: lambda(t) = beta S(t)/N",
            "At each exposure, a user can share, saturate, or move forward",
            "Condition: 0 <= rho_i, sigma_i <= 1 and rho_i + sigma_i <= 1",
        ],
        top=1.45,
        height=4.6,
        fill=(255, 247, 235),
    )

    # Slide 4: equations
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "3. Governing Equations", "The key point is how the exposure chain feeds into sharing")
    add_footer(slide)
    eq_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(7.7), Inches(4.8))
    tf = eq_box.text_frame
    lines = [
        ("lambda(t) = beta S(t) / N", 24, True),
        ("dI0/dt = - lambda I0", 18, False),
        ("dI1/dt = (1-rho0-sigma0) lambda I0 - lambda I1", 18, False),
        ("dI2/dt = (1-rho1-sigma1) lambda I1 - (rho2+sigma2) lambda I2", 18, False),
        ("dS/dt = rho0 lambda I0 + rho1 lambda I1 + rho2 lambda I2 - gamma S", 18, False),
        ("dR/dt = sigma0 lambda I0 + sigma1 lambda I1 + sigma2 lambda I2 + gamma S", 18, False),
    ]
    for idx, (text, size, bold) in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor(32, 39, 49)
        p.space_after = Pt(10)
    add_side_panel(
        slide,
        "How to explain I0 -> I1",
        [
            "Users in I0 are exposed at rate lambda I0.",
            "If they neither share nor saturate after that exposure, they move to I1.",
            "That is why the inflow into I1 is (1-rho0-sigma0) lambda I0.",
        ],
        top=1.55,
        height=2.45,
        fill=(232, 240, 254),
    )
    add_side_panel(
        slide,
        "Interpretation",
        [
            "beta makes exposure stronger",
            "gamma removes current sharers",
            "rho_i promotes growth of S",
            "sigma_i removes users from further spreading",
        ],
        top=4.25,
        height=1.8,
        fill=(232, 245, 236),
    )

    # Slide 5: equilibrium and stability
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "4. Equilibria and Local Stability", "Professor feedback addressed: equilibrium first, then Jacobian")
    add_footer(slide)
    add_bullets(
        slide,
        [
            "Disease-free equilibrium: (I0, I1, I2, S) = (N, 0, 0, 0), with R = 0.",
            "Interpretation: nobody is sharing, so no new exposures occur.",
            "Burned-out equilibrium: (I0, I1, I2, S) = (0, 0, 0, 0), with R = N.",
            "Interpretation: the content has fully died out and everyone is in R.",
        ],
        width=7.7,
        font_size=19,
    )
    jac_box = slide.shapes.add_textbox(Inches(0.85), Inches(4.9), Inches(7.8), Inches(0.8))
    pj = jac_box.text_frame.paragraphs[0]
    pj.text = "J(E0) = [[0,0,0,-beta], [0,0,0,(1-rho0-sigma0)beta], [0,0,0,0], [0,0,0,rho0 beta-gamma]]"
    pj.font.size = Pt(16)
    pj.font.color.rgb = RGBColor(32, 39, 49)
    add_side_panel(
        slide,
        "Crucial clarification",
        [
            "At the disease-free equilibrium, I0 = N.",
            "This is not the same as the simulation start.",
            "In simulations, I use a seed such as S(0)=15, so I0(0)=985.",
        ],
        top=1.45,
        height=2.15,
        fill=(255, 247, 235),
    )
    add_side_panel(
        slide,
        "Sign checks",
        [
            "dI0/dt = -lambda I0 <= 0, so I0 never increases.",
            "dR/dt >= 0, so R never decreases.",
            "The key local sign is rho0 beta - gamma.",
            "Negative: tiny sharing dies out. Positive: tiny sharing can grow initially.",
        ],
        top=3.95,
        height=2.15,
        fill=(232, 240, 254),
    )

    # Slide 6: simulation setup and outcomes
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "5. Simulation Setup and Main Outcomes", "Three scenarios: viral, moderate, and fizzle")
    add_footer(slide)
    setup_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(3.6), Inches(2.1))
    tf = setup_box.text_frame
    for idx, text in enumerate(
        [
            "m = 2, N = 1000",
            "Initial seed: S(0) = 15",
            "I0(0) = 985, I1(0) = I2(0) = 0, R(0) = 0",
            "Integrated over t in [0, 120]",
        ]
    ):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = text
        p.font.size = Pt(19)
        p.font.color.rgb = RGBColor(32, 39, 49)
        p.space_after = Pt(10)
    result_rows = [
        ["Scenario", "Peak S(t)", "Time to peak", "R(120)"],
        ["Viral", f"{stats['viral_peak']:.2f}", f"{stats['viral_peak_t']:.2f}", f"{stats['viral_R']:.2f}"],
        ["Moderate", f"{stats['moderate_peak']:.2f}", f"{stats['moderate_peak_t']:.2f}", f"{stats['moderate_R']:.2f}"],
        ["Fizzle", f"{stats['fizzle_peak']:.2f}", f"{stats['fizzle_peak_t']:.2f}", f"{stats['fizzle_R']:.2f}"],
    ]
    add_small_table(slide, result_rows, left=0.8, top=3.1, width=7.8, height=2.2)
    add_side_panel(
        slide,
        "Result summary",
        [
            "Viral reaches the largest and earliest large outbreak.",
            "Moderate still spreads, but more slowly and less strongly.",
            "Fizzle stays close to the initial seed and never develops into a cascade.",
        ],
        fill=(240, 244, 248),
    )

    # Slide 7: scenario comparison figure
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "6. Comparison of Sharer Dynamics", "The main result figure for peer interpretation")
    add_footer(slide)
    slide.shapes.add_picture(str(assets["comparison"]), Inches(0.8), Inches(1.4), width=Inches(7.9))
    add_caption(slide, "Figure: Active sharers S(t) for viral, moderate, and fizzle scenarios.")
    add_side_panel(
        slide,
        "Interpretation",
        [
            "The three scenarios separate clearly.",
            "Higher effective spread creates a larger and earlier peak.",
            "Weak spread leaves the system close to the initial condition.",
            "This figure directly supports the project's main comparative result.",
        ],
        fill=(232, 240, 254),
    )

    # Slide 8: exposure classes
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "7. Exposure-Class Dynamics", "Why the multi-exposure structure matters")
    add_footer(slide)
    slide.shapes.add_picture(str(assets["exposure"]), Inches(0.8), Inches(1.4), width=Inches(7.9))
    add_caption(slide, "Figure: Viral scenario showing I0, I1, I2, S, and R over time.")
    add_side_panel(
        slide,
        "Interpretation",
        [
            "Users leave I0 and move through higher exposure classes.",
            "Some eventually join S, while others end in R.",
            "This makes the reinforcement mechanism visible in a way a one-step model cannot.",
        ],
        fill=(255, 247, 235),
    )

    # Slide 9: sensitivity
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "8. Sensitivity to Gamma", "How sharer decay changes the peak")
    add_footer(slide)
    slide.shapes.add_picture(str(assets["sensitivity"]), Inches(0.8), Inches(1.4), width=Inches(7.9))
    add_caption(slide, "Figure: Peak active sharers as gamma varies, with other parameters fixed.")
    add_side_panel(
        slide,
        "Interpretation",
        [
            "Gamma measures how quickly current sharers stop being active.",
            "As gamma increases, the peak drops sharply.",
            "In plain language: if attention fades faster, the post is much less likely to spread widely.",
        ],
        fill=(232, 245, 236),
    )

    # Slide 10: remaining work
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "9. What Still Needs To Be Completed", "Included explicitly to match the presentation guidelines")
    add_footer(slide)
    add_bullets(
        slide,
        [
            "Refine the equilibrium and stability discussion in the written report.",
            "Present the Jacobian analysis more cleanly and tie it more directly to the simulations.",
            "Extend sensitivity analysis for gamma, rho_i, and sigma_i.",
            "Strengthen the written discussion of figures and improve figure captions.",
            "Tighten the interpretation connecting the mathematics to social-media behaviour.",
        ],
        width=7.9,
        font_size=21,
    )
    add_side_panel(
        slide,
        "Current status",
        [
            "The model formulation is complete.",
            "The main simulations and figures are complete.",
            "Most remaining work is analysis, explanation, and polishing for the final report.",
        ],
        fill=(240, 244, 248),
    )

    # Slide 11: takeaway
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide)
    add_header_bar(slide, "10. Conclusion and Q&A", "A concise ending that leaves time for discussion")
    add_footer(slide)
    add_bullets(
        slide,
        [
            "Repeated exposure matters for repost behaviour.",
            "A multi-exposure ODE model can produce viral, moderate, and fizzle outcomes within one framework.",
            "The balance among beta, gamma, rho_i, and sigma_i determines whether content grows or dies out.",
            "The model is simple enough to analyze, but rich enough to capture reinforcement before sharing.",
        ],
        width=7.9,
        font_size=20,
    )
    add_side_panel(
        slide,
        "Likely Q&A themes",
        [
            "Why not use a network model?",
            "Why is I0 = N at equilibrium but not in simulations?",
            "What do rho_i and sigma_i mean in social-media terms?",
            "What does the Jacobian tell us, and what does it not tell us?",
        ],
        fill=(232, 240, 254),
    )

    prs.save(PPTX_PATH)

    notes = """# Project Presentation Speaker Notes

## Overall timing

- Aim for about 11 to 12 minutes of speaking.
- Leave the final 2 to 3 minutes for questions.
- Do not rush the equations slide; that is the slide most likely to generate questions.

## Core clarification you must say clearly

There are two different ideas that should not be mixed up:

- The disease-free equilibrium is the exact no-spread state, so `I0 = N` and all other compartments are `0`.
- The simulations do not start at that equilibrium. They start with a small seed of sharers, so `S(0) > 0` and therefore `I0(0) = N - S(0)`.

So if someone asks whether `I0` should be `N` or `0`, the correct answer is:

- `I0 = N` at the disease-free equilibrium
- `I0(0) = N - S(0)` in the simulation setup

## Slide-by-slide notes

### Slide 1: Title

Open simply: this project studies how one piece of content spreads on social media when repeated exposure matters. Mention that the presentation will cover the model, the main results, and what remains to finish for the final report.

### Slide 2: Motivation and background

Explain that on social media, seeing a post once is often not enough for someone to repost it. Some users share quickly, some need reinforcement, and some lose interest. That is the motivation for tracking exposure classes rather than using a single susceptible or ignorant compartment.

You can say that the project is related to rumor models, threshold models, and complex contagion, but your contribution is to keep the model in a simple ODE form while still allowing repeated exposure to matter.

### Slide 3: Compartments, parameters, and assumptions

Walk through the table slowly enough for the audience to see the structure of the model. Explain that `rho_i` and `sigma_i` were important comments from the professor, so you now state them clearly:

- `rho_i` is the probability of sharing after exposure level `i`
- `sigma_i` is the probability of saturating, or losing interest, after exposure level `i`

Then mention the condition `rho_i + sigma_i <= 1`, which ensures the remaining probability corresponds to moving forward in the exposure chain.

### Slide 4: Governing equations

This is the main mechanics slide. First define `lambda(t) = beta S(t)/N`. Then explain that:

- `beta` controls how strongly sharers expose others
- `gamma` controls how quickly current sharers stop being active

For the `I0` to `I1` step, use this explanation almost exactly:

`Users in I0 are exposed at rate lambda I0. If they do not share and do not saturate after that exposure, they move into I1. That is why the inflow into I1 is (1-rho0-sigma0) lambda I0.`

That phrasing is simple and mathematically correct.

### Slide 5: Equilibria and local stability

Present the disease-free equilibrium first, then the burned-out equilibrium, and only after that mention the Jacobian. This matches the order your professor wanted.

When you get to the sign discussion, say:

`There are two useful sign checks in this model. First, dI0/dt = -lambda I0 is always nonpositive, so I0 never increases. Second, dR/dt is always nonnegative, so R never decreases. For local stability near the disease-free equilibrium, the key sign is rho0 beta minus gamma.`

Then explain:

- if `rho0 beta - gamma < 0`, tiny sharing dies out initially
- if `rho0 beta - gamma > 0`, tiny sharing can grow initially

Finish by saying that this Jacobian result is local, so the simulations are needed to understand the broader nonlinear behaviour.

### Slide 6: Simulation setup and main outcomes

State the simulation setup clearly: `m = 2`, `N = 1000`, and an initial seed `S(0) = 15`. Then summarize the table rather than reading every number.

The main message is:

- the viral scenario produces the largest outbreak
- the moderate scenario spreads, but more slowly and less strongly
- the fizzle scenario stays close to the seed and does not develop into a cascade

### Slide 7: Comparison of sharer dynamics

Use the figure to tell a story. Point first to the viral curve, then the moderate curve, then the fizzle curve. Emphasize that this is the clearest summary of the project's numerical result because the three regimes are visually distinct.

### Slide 8: Exposure-class dynamics

Explain that this figure justifies the multi-exposure model itself. It shows users leaving `I0`, passing through the exposure classes, and then eventually entering either `S` or `R`. This is what a one-step model would hide.

### Slide 9: Sensitivity to gamma

Explain `gamma` in plain language: it measures how quickly current sharers lose interest and stop participating in the spread. Then use the graph to show that when `gamma` increases, the peak drops substantially. This gives a clear interpretation of one of the main parameters.

### Slide 10: What still needs to be completed

Be direct and professional here. Say that the model and the main numerical results are already complete, and the remaining work is mostly strengthening the analysis and polishing the report. That sounds honest and academically mature.

### Slide 11: Conclusion and Q&A

End with three points:

- repeated exposure matters
- the model can reproduce very different spreading regimes
- the balance among `beta`, `gamma`, `rho_i`, and `sigma_i` determines whether the post grows or dies out

Then pause and invite questions.

## Likely questions and polished answers

### Why not use a network model?

A network model would be a natural extension, but for this project I wanted a model that remains analytically tractable while still capturing repeated exposure. The ODE framework gives that balance.

### Why is `I0 = N` at equilibrium but not in the simulation?

The equilibrium describes the exact no-spread resting state. The simulation is intentionally started away from that state by introducing a small seed of sharers, so `I0(0)` must be less than `N`.

### What do `rho_i` and `sigma_i` represent in practice?

`rho_i` represents the chance that a user reposts after exposure level `i`, while `sigma_i` represents the chance that the user loses interest and stops participating after that exposure level.

### What does the Jacobian tell us?

It tells us the local behaviour near the disease-free equilibrium. In particular, it identifies the sign condition that determines whether a very small amount of sharing initially grows or dies out.

### Why are simulations still necessary if you have the Jacobian?

Because the Jacobian is only a local linear approximation. The full model includes nonlinear flow through multiple exposure classes, so simulations are needed to understand the larger cascade behaviour.
"""
    NOTES_PATH.write_text(notes)


if __name__ == "__main__":
    build_ppt()
